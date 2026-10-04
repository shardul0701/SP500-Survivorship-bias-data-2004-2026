"""Daily backstop: diff the PIT YAML against the index's CURRENT official member list.

Why this exists
---------------
The release crawler can only add what it can find and parse. In 2026 it missed
five Nasdaq-100 events, and nothing reported any of them:

* Lumentum-for-CoStar (05-18) -- release missing from the IR archive listing,
  and "CoStar Group, Inc." broke the sentence-scoped parser;
* Moderna-for-WBD (10-09) -- "( Nasdaq : MRNA)" defeated the ticker pattern;
* the HONA spin-off (06-29), the EA take-private (08-04) and the KHC move to
  NYSE (09-14) -- none of which is ever announced as a press release.

Every one of them is visible in a single comparison: what the index provider
says the members are TODAY versus what the YAML says they are today. A
one-for-one swap keeps the count at 101, so a count check cannot see it; a
name-level diff can. This module is that diff. It is deliberately independent
of the crawler -- it shares no parsing with it -- so a blind crawler cannot
also blind the check.

Boundary tolerance: a provider can publish a change a day early or late
relative to the effective date in the YAML. A name is reported only if it
disagrees with the YAML on EVERY weekday in [as_of - 1, as_of + 2].

Self-contained on purpose (PyYAML + requests, openpyxl for S&P) so the same
file serves both PIT repositories.
"""

from __future__ import annotations

import io
import json
import re
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import requests
import yaml

ROOT = Path(__file__).resolve().parents[1]
METADATA_DIR = ROOT / "metadata"
STATE_FILE = METADATA_DIR / "live_check_state.json"

PROFILES = {
    "nq100": {
        "index_name": "Nasdaq-100",
        "data_dir": ROOT / "src" / "nasdaq_100_ticker_history",
        "filename": "n100-ticker-changes-{year}.yaml",
        "min_members": 95,
        "max_members": 110,
        "official": "nasdaq_api",
        "advisory": ["upstream_n100tickers"],
    },
    "sp500": {
        "index_name": "S&P 500",
        "data_dir": ROOT / "src" / "sp500_ticker_history",
        "filename": "sp500-ticker-changes-{year}.yaml",
        "min_members": 490,
        "max_members": 520,
        "official": "ssga_spy",
        "advisory": [],
    },
}

BROWSER_UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/128.0 Safari/537.36"
)
TICKER_SHAPE = re.compile(r"^[A-Z][A-Z0-9.\-/ ]{0,11}$")


@dataclass
class SourceResult:
    name: str
    url: str
    role: str  # "official" gates the status; "advisory" is reported only
    reachable: bool = False
    as_of: str | None = None
    count: int = 0
    error: str | None = None
    missing_from_pit: list[str] = field(default_factory=list)
    extra_in_pit: list[str] = field(default_factory=list)
    window: list[str] = field(default_factory=list)


# --------------------------------------------------------------------------- #
# Ticker comparison
# --------------------------------------------------------------------------- #
def load_aliases() -> dict[str, str]:
    path = METADATA_DIR / "ticker_mapping.yaml"
    if not path.exists():
        return {}
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    return {str(k).upper(): str(v).upper() for k, v in (data.get("aliases") or {}).items()}


def compare_key(ticker: str, aliases: dict[str, str]) -> str:
    """Share-class punctuation differs by publisher (BRK.B / BRK-B / BRK/B /
    "BRK B"); strip it so only a genuinely different symbol counts."""
    t = str(ticker).strip().upper()
    t = aliases.get(t, t)
    return re.sub(r"[.\-/ ]", "", t)


# --------------------------------------------------------------------------- #
# PIT membership from this repository's YAML
# --------------------------------------------------------------------------- #
def members_from_yaml(data: dict, on: date) -> set[str]:
    members = set(data.get("tickers_on_Jan_1") or [])
    for key, entry in sorted((data.get("changes") or {}).items(), key=lambda kv: str(kv[0])):
        if str(key) > on.isoformat():
            break
        entry = entry or {}
        members -= set(entry.get("difference") or [])
        members |= set(entry.get("union") or [])
    return members


def pit_members(index: str, on: date) -> set[str]:
    prof = PROFILES[index]
    path = prof["data_dir"] / prof["filename"].format(year=on.year)
    if path.exists():
        return members_from_yaml(yaml.safe_load(path.read_text(encoding="utf-8")) or {}, on)
    # The window around Dec 31 reaches into a year whose file rollover_year
    # has not created yet; its membership is last year's final state.
    prev = prof["data_dir"] / prof["filename"].format(year=on.year - 1)
    if not prev.exists():
        raise FileNotFoundError(f"no PIT YAML for {on.year} or {on.year - 1}")
    return members_from_yaml(
        yaml.safe_load(prev.read_text(encoding="utf-8")) or {}, date(on.year - 1, 12, 31)
    )


def weekday_window(as_of: date, back: int = 1, forward: int = 2) -> list[date]:
    def step(d: date, n: int) -> date:
        sign = 1 if n > 0 else -1
        for _ in range(abs(n)):
            d += timedelta(days=sign)
            while d.weekday() >= 5:
                d += timedelta(days=sign)
        return d

    out = [step(as_of, -i) for i in range(back, 0, -1)] + [as_of]
    out += [step(as_of, i) for i in range(1, forward + 1)]
    return out


def persistent_diff(
    members_on, live: set[str], as_of: date, aliases: dict[str, str]
) -> tuple[list[str], list[str], list[str]]:
    """Names that disagree with the YAML on every day of the tolerance window.

    members_on(day) -> set of YAML members that day.
    """
    window = weekday_window(as_of)
    live_keys = {compare_key(t, aliases): t for t in live}
    missing: set[str] | None = None
    extra: set[str] | None = None
    for day in window:
        pit = {compare_key(t, aliases): t for t in members_on(day)}
        day_missing = {live_keys[k] for k in live_keys.keys() - pit.keys()}
        day_extra = {pit[k] for k in pit.keys() - live_keys.keys()}
        missing = day_missing if missing is None else missing & day_missing
        extra = day_extra if extra is None else extra & day_extra
    return sorted(missing or ()), sorted(extra or ()), [d.isoformat() for d in window]


# --------------------------------------------------------------------------- #
# Live sources
# --------------------------------------------------------------------------- #
def fetch_nasdaq_api() -> tuple[set[str], date, str]:
    url = "https://api.nasdaq.com/api/quote/list-type/nasdaq100"
    resp = requests.get(
        url,
        headers={
            "User-Agent": BROWSER_UA,
            "Accept": "application/json, text/plain, */*",
            "Origin": "https://www.nasdaq.com",
            "Referer": "https://www.nasdaq.com/",
        },
        timeout=30,
    )
    resp.raise_for_status()
    payload = (resp.json() or {}).get("data") or {}
    rows = ((payload.get("data") or {}).get("rows")) or []
    tickers = {str(r["symbol"]).strip().upper() for r in rows if r.get("symbol")}
    stamp = payload.get("date") or (payload.get("data") or {}).get("asOf")
    as_of = datetime.strptime(stamp.strip(), "%b %d, %Y").date() if stamp else date.today()
    return tickers, as_of, url


def fetch_upstream_n100tickers() -> tuple[set[str], date, str]:
    """Community cross-check (jmccarrell/n100tickers, this dataset's 2015+
    base). Not official and sometimes days late, so it never gates the status;
    it is reported so a disagreement with BOTH sources is easy to spot."""
    as_of = date.today()
    url = (
        "https://raw.githubusercontent.com/jmccarrell/n100tickers/main/src/"
        f"nasdaq_100_ticker_history/n100-ticker-changes-{as_of.year}.yaml"
    )
    resp = requests.get(url, timeout=30)
    resp.raise_for_status()
    return members_from_yaml(yaml.safe_load(resp.text) or {}, as_of), as_of, url


def fetch_ssga_spy() -> tuple[set[str], date, str]:
    """SPY's daily holdings file: the S&P 500 as State Street holds it today."""
    from openpyxl import load_workbook

    url = (
        "https://www.ssga.com/us/en/intermediary/library-content/products/"
        "fund-data/etfs/us/holdings-daily-us-en-spy.xlsx"
    )
    resp = requests.get(url, headers={"User-Agent": BROWSER_UA}, timeout=60)
    resp.raise_for_status()
    sheet = load_workbook(io.BytesIO(resp.content), read_only=True, data_only=True).active
    rows = [list(r) for r in sheet.iter_rows(values_only=True)]
    as_of = None
    header = None
    for i, row in enumerate(rows):
        for cell in row:
            m = re.search(r"As of\s+(\d{1,2}-[A-Za-z]{3}-\d{4})", str(cell or ""))
            if m and as_of is None:
                as_of = datetime.strptime(m.group(1), "%d-%b-%Y").date()
        if row and str(row[0] or "").strip() == "Name":
            header = i
            break
    if header is None:
        raise ValueError("SPY holdings: no 'Name' header row")
    cols = [str(c or "").strip() for c in rows[header]]
    t_col, n_col = cols.index("Ticker"), cols.index("Name")
    tickers = set()
    for row in rows[header + 1 :]:
        if len(row) <= t_col or row[t_col] is None:
            continue
        ticker = str(row[t_col]).strip().upper()
        name = str(row[n_col] or "").upper()
        if not TICKER_SHAPE.match(ticker) or "CASH" in ticker or "CASH" in name:
            continue
        tickers.add(ticker)
    return tickers, as_of or date.today(), url


SOURCES = {
    "nasdaq_api": fetch_nasdaq_api,
    "upstream_n100tickers": fetch_upstream_n100tickers,
    "ssga_spy": fetch_ssga_spy,
}


# --------------------------------------------------------------------------- #
# Run
# --------------------------------------------------------------------------- #
def run_check(index: str) -> dict:
    prof = PROFILES[index]
    aliases = load_aliases()
    results: list[SourceResult] = []
    plan = [(prof["official"], "official")] + [(n, "advisory") for n in prof["advisory"]]
    for name, role in plan:
        res = SourceResult(name=name, url="", role=role)
        try:
            live, as_of, url = SOURCES[name]()
            res.url, res.as_of, res.count = url, as_of.isoformat(), len(live)
            if not prof["min_members"] <= len(live) <= prof["max_members"]:
                raise ValueError(
                    f"{len(live)} members is outside the plausible "
                    f"{prof['min_members']}-{prof['max_members']} range -- the "
                    f"source format probably changed"
                )
            res.missing_from_pit, res.extra_in_pit, res.window = persistent_diff(
                lambda day: pit_members(index, day), live, as_of, aliases
            )
            res.reachable = True
        except Exception as exc:  # recorded, never swallowed: see status below
            res.error = f"{type(exc).__name__}: {exc}"[:300]
        results.append(res)

    official = results[0]
    if not official.reachable:
        status = "source_error"
    elif official.missing_from_pit or official.extra_in_pit:
        status = "mismatch"
    else:
        status = "pass"
    today = date.today()
    state = {
        "index_name": prof["index_name"],
        "checked_at": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
        "status": status,
        "pit_count_today": len(pit_members(index, today)),
        "official_source": official.name,
        "sources": [asdict(r) for r in results],
    }
    METADATA_DIR.mkdir(parents=True, exist_ok=True)
    previous = load_state()
    if status == "pass":
        state["last_pass_at"] = state["checked_at"]
    else:
        state["last_pass_at"] = (previous or {}).get("last_pass_at")
    STATE_FILE.write_text(json.dumps(state, indent=2) + "\n", encoding="utf-8")
    return state


def load_state() -> dict | None:
    if not STATE_FILE.exists():
        return None
    return json.loads(STATE_FILE.read_text(encoding="utf-8"))


# A single unreachable fetch is usually the provider, not us; a name-level
# disagreement never is. Mismatches alert immediately; an unreachable official
# source alerts once it has gone this long without a passing check.
SOURCE_ERROR_GRACE_DAYS = 3


def needs_attention(state: dict, now: datetime | None = None) -> bool:
    if state["status"] == "mismatch":
        return True
    if state["status"] != "source_error":
        return False
    last = state.get("last_pass_at")
    if not last:
        return True
    now = now or datetime.now(timezone.utc)
    return (now - datetime.fromisoformat(last)).days >= SOURCE_ERROR_GRACE_DAYS


def render_issue(state: dict) -> tuple[str, str]:
    """Title and Markdown body for the live-check alert."""
    name = state["index_name"]
    off = state["sources"][0]
    if state["status"] == "mismatch":
        title = f"[live-check] {name} PIT disagrees with the official current member list"
    else:
        title = f"[live-check] {name} official member list could not be checked"
    lines = [
        f"**Status:** `{state['status']}` · checked {state['checked_at']} · "
        f"PIT members today: {state['pit_count_today']}",
        "",
        f"Official source `{off['name']}` ({off['url'] or 'n/a'}) as of "
        f"{off['as_of'] or 'n/a'}, {off['count']} members.",
        "",
    ]
    if off["error"]:
        lines += [f"Fetch error: `{off['error']}`", ""]
    if off["missing_from_pit"] or off["extra_in_pit"]:
        lines += [
            "| | Tickers |",
            "|---|---|",
            f"| In the official list, **missing from the YAML** | "
            f"{', '.join(off['missing_from_pit']) or '-'} |",
            f"| In the YAML, **not in the official list** | "
            f"{', '.join(off['extra_in_pit']) or '-'} |",
            "",
            f"A name is listed only if it disagrees on every weekday in "
            f"{off['window'][0]} .. {off['window'][-1]}, so a change published a "
            f"day early or late is not reported.",
            "",
        ]
    for adv in state["sources"][1:]:
        if adv["reachable"]:
            lines.append(
                f"Advisory `{adv['name']}` (as of {adv['as_of']}): missing from YAML "
                f"{adv['missing_from_pit'] or '-'}, extra in YAML {adv['extra_in_pit'] or '-'}."
            )
        else:
            lines.append(f"Advisory `{adv['name']}` unreachable: `{adv['error']}`.")
    lines += [
        "",
        "**To fix:** add a dated change entry for each name to the yearly YAML. "
        "Use `source_url` when an official release exists; spin-offs, take-privates "
        "and exchange transfers have none, so record `evidence_url` + `evidence_note` "
        "instead. The issue closes itself on the first run that passes.",
    ]
    return title, "\n".join(lines)
